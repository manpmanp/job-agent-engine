# Job-search agent engine
# File:    engine/webpage.py
# Purpose: Renders the weekly page with Accept/Reject buttons.
# Author:  manpmanp <manp.node@gmail.com>
# Generated with Claude (Anthropic), under the author's direction and review.
# License: MIT (see LICENSE)
# Version: 0.3.0
# Date:    2026-10-28

"""
Weekly page: renders the Active sheet of the tracker as one standalone HTML
file (no server, no build step) with Accept/Reject buttons.

- Accept queues a posting for tailored documents (this click is the
  per-posting approval).
- Reject asks for one or more reasons (configurable in
  search_config.rejection_reasons).
- Decisions live in the browser (localStorage 'jobAgentDecisions:<home id>',
  one key per agent home, see home_id()) until
  "Sync decisions" writes decisions_sync.json into the agent's state folder.
  './run.sh sync' then applies them (see engine/decisions.py).
- Rows are grouped client-side: new this week, needs a decision, accepted,
  rejected.
"""
import html
import json
import os
from datetime import datetime, date

import openpyxl

from .tracker import to_date, seniority_bucket

PAGE_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{page_title}</title>
<style>
  :root {{
    --bg:#0f1115; --panel:#171a21; --panel2:#1e222b; --border:#2a2f3a;
    --text:#e6e9ef; --muted:#9aa4b2; --accent:#4ade80; --link:#60a5fa;
    --senior-bg:#3a2a12; --senior-fg:#f5b662; --grad-bg:#3a1414; --grad-fg:#f28b8b;
    --good-bg:#12301e; --good-fg:#5fd88a; --other-bg:#1e222b; --other-fg:#9aa4b2;
    --warn-bg:#3a1414; --warn-fg:#ff8080;
    --accept-bg:#12301e; --accept-fg:#5fd88a; --accept-border:#245c39;
    --reject-bg:#3a1414; --reject-fg:#ff8080; --reject-border:#5c2424;
    --new-bg:#1a2540; --new-fg:#8ab4ff;
  }}
  * {{ box-sizing:border-box; }}
  body {{
    margin:0; padding:0; background:var(--bg); color:var(--text);
    font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif;
  }}
  header {{
    padding:28px 32px 18px; border-bottom:1px solid var(--border);
    background:var(--panel);
  }}
  header h1 {{ margin:0 0 6px; font-size:20px; font-weight:600; }}
  header p {{ margin:0; color:var(--muted); font-size:13px; }}
  .meta {{ display:flex; gap:22px; margin-top:14px; flex-wrap:wrap; }}
  .meta div {{ font-size:12px; color:var(--muted); }}
  .meta strong {{ color:var(--text); font-size:15px; display:block; }}
  .controls {{
    padding:14px 32px; display:flex; gap:10px; flex-wrap:wrap; align-items:center;
    background:var(--panel); border-bottom:1px solid var(--border);
  }}
  .controls input[type=text] {{
    background:var(--panel2); border:1px solid var(--border); color:var(--text);
    padding:7px 10px; border-radius:6px; font-size:13px; width:220px;
  }}
  .controls label {{ font-size:12px; color:var(--muted); display:flex; gap:6px; align-items:center; }}
  .tabs {{ display:flex; gap:6px; }}
  .tab {{
    background:var(--panel2); border:1px solid var(--border); color:var(--muted);
    padding:6px 12px; border-radius:6px; font-size:12px; cursor:pointer;
  }}
  .tab.active {{ color:var(--text); border-color:var(--link); }}
  .btn {{
    border:1px solid var(--border); background:var(--panel2); color:var(--text);
    padding:6px 12px; border-radius:6px; font-size:12px; cursor:pointer;
  }}
  .btn:hover {{ border-color:var(--link); }}
  .wrap {{ padding:20px 32px 48px; overflow-x:auto; }}
  table {{ border-collapse:collapse; width:100%; min-width:1250px; }}
  th, td {{
    padding:10px 12px; border-bottom:1px solid var(--border); font-size:13px;
    text-align:left; vertical-align:top;
  }}
  th {{
    background:var(--panel2); color:var(--muted); font-weight:600;
    font-size:11px; text-transform:uppercase; letter-spacing:.04em;
    position:sticky; top:0;
  }}
  tr:hover td {{ background:#1a1e27; }}
  .badge {{
    display:inline-block; padding:3px 8px; border-radius:5px; font-size:11px;
    font-weight:600; white-space:nowrap;
  }}
  .sen-senior {{ background:var(--senior-bg); color:var(--senior-fg); }}
  .sen-grad {{ background:var(--grad-bg); color:var(--grad-fg); }}
  .sen-good {{ background:var(--good-bg); color:var(--good-fg); }}
  .sen-other {{ background:var(--other-bg); color:var(--other-fg); }}
  .action {{ color:var(--warn-fg); font-weight:600; font-size:12px; }}
  a.posting-link {{ color:var(--link); text-decoration:none; }}
  a.posting-link:hover {{ text-decoration:underline; }}
  .employer {{ font-weight:600; }}
  .muted {{ color:var(--muted); }}
  footer {{ padding:20px 32px 40px; color:var(--muted); font-size:12px; }}
  .hidden {{ display:none; }}

  .group-header-row td {{
    background:var(--panel2); color:var(--muted); font-weight:700;
    font-size:11px; text-transform:uppercase; letter-spacing:.06em;
    padding:9px 12px; border-bottom:1px solid var(--border);
    border-top:2px solid var(--border);
  }}
  .group-header-row.empty-group {{ display:none; }}
  .group-count {{ font-weight:700; color:var(--text); margin-left:4px; }}
  tr[data-key].is-new td {{ box-shadow: inset 3px 0 0 var(--new-fg); }}
  .new-tag {{
    display:inline-block; background:var(--new-bg); color:var(--new-fg);
    font-size:10px; font-weight:700; padding:2px 6px; border-radius:4px;
    margin-left:6px; text-transform:uppercase; letter-spacing:.03em; vertical-align:middle;
  }}

  .decision-col {{ min-width:190px; }}
  .decision-buttons {{ display:flex; gap:6px; }}
  .btn-accept, .btn-reject {{
    border-radius:6px; padding:6px 10px; font-size:12px; font-weight:600;
    cursor:pointer; border:1px solid;
  }}
  .btn-accept {{ background:var(--accept-bg); color:var(--accept-fg); border-color:var(--accept-border); }}
  .btn-accept:hover {{ filter:brightness(1.2); }}
  .btn-reject {{ background:var(--reject-bg); color:var(--reject-fg); border-color:var(--reject-border); }}
  .btn-reject:hover {{ filter:brightness(1.2); }}
  .decision-badge {{
    display:inline-block; padding:4px 9px; border-radius:6px; font-size:11px;
    font-weight:600; line-height:1.4;
  }}
  .decision-badge.accepted {{ background:var(--accept-bg); color:var(--accept-fg); }}
  .decision-badge.rejected {{ background:var(--reject-bg); color:var(--reject-fg); }}
  .decision-reason {{ display:block; font-weight:400; font-style:italic; opacity:.85; }}
  .undo-link {{
    display:block; margin-top:4px; font-size:11px; color:var(--muted);
    cursor:pointer; text-decoration:underline;
  }}
  .pending-note {{ display:block; font-size:10px; color:var(--muted); margin-top:2px; }}

  .reason-picker {{
    background:var(--panel2); border:1px solid var(--reject-border); border-radius:8px;
    padding:10px; margin-top:6px; width:260px;
  }}
  .reason-picker .r-title {{ font-size:11px; color:var(--muted); margin-bottom:6px; text-transform:uppercase; letter-spacing:.04em; }}
  .reason-picker label {{
    display:flex; gap:6px; align-items:flex-start; font-size:12px; padding:3px 0; cursor:pointer;
  }}
  .reason-picker input[type=text] {{
    width:100%; margin-top:6px; background:var(--panel); border:1px solid var(--border);
    color:var(--text); border-radius:5px; padding:5px 7px; font-size:12px;
  }}
  .reason-picker .r-actions {{ display:flex; gap:6px; margin-top:8px; }}
  .reason-picker .r-actions button {{
    flex:1; padding:6px; border-radius:5px; font-size:12px; cursor:pointer; border:1px solid var(--border);
    background:var(--panel); color:var(--text);
  }}
  .reason-picker .r-actions button.confirm {{ background:var(--reject-bg); color:var(--reject-fg); border-color:var(--reject-border); }}
</style>
</head>
<body>
<header>
  <h1>{page_title_html}</h1>
  <p>Auto-generated from the tracker. Accept queues a posting for a tailored CV/cover letter; Reject records why, for reference in future runs. Postings are grouped below: new this week, needs a decision, accepted, rejected.</p>
  <div class="meta">
    <div><strong>{active_count}</strong>active</div>
    <div><strong>{urgent_count}</strong>flagged (act soon)</div>
    <div><strong id="newCount">0</strong>new this week</div>
    <div><strong id="acceptedCount">0</strong>accepted</div>
    <div><strong id="rejectedCount">0</strong>rejected</div>
    <div><strong>{generated_at}</strong>last generated{weekly_note_html}</div>
  </div>
</header>
<div class="controls">
  <input type="text" id="filterBox" placeholder="Filter by employer, title, location..." onkeyup="filterRows()">
  <div class="tabs" id="decisionTabs">
    <div class="tab active" data-tab="all" onclick="setTab('all')">All</div>
    <div class="tab" data-tab="undecided" onclick="setTab('undecided')">Undecided</div>
    <div class="tab" data-tab="accepted" onclick="setTab('accepted')">Accepted</div>
    <div class="tab" data-tab="rejected" onclick="setTab('rejected')">Rejected</div>
  </div>
  <label><input type="checkbox" id="onlyFlagged" onchange="filterRows()"> Only flagged (act soon)</label>
  <label><input type="checkbox" id="hideSenior" onchange="filterRows()"> Hide flagged as too senior</label>
  <button class="btn" onclick="syncDecisions()" title="Saves decisions_sync.json into {sync_hint_attr} (one-time folder pick), or downloads it if your browser does not support that. The agent reads it on its next sync.">Sync decisions</button>
  <span id="syncStatus" class="muted" style="font-size:11px;"></span>
</div>
<div class="wrap">
<table id="jobsTable">
  <thead>
    <tr>
      <th class="decision-col">Decision</th>
      <th>Employer</th>
      <th>Title</th>
      <th>Seniority</th>
      <th>Location</th>
      <th>Mode</th>
      <th>Announced</th>
      <th>Days</th>
      <th>Category</th>
      <th>Skills</th>
      <th>Link</th>
      <th>Status</th>
    </tr>
  </thead>
  <tbody id="group-new">
    <tr class="group-header-row"><td colspan="12">&#x1F195; New this week <span class="group-count" id="count-new">0</span></td></tr>
  </tbody>
  <tbody id="group-undecided">
    <tr class="group-header-row"><td colspan="12">&#x1F550; Needs a decision <span class="group-count" id="count-undecided">0</span></td></tr>
  </tbody>
  <tbody id="group-accepted">
    <tr class="group-header-row"><td colspan="12">&#x2705; Accepted <span class="group-count" id="count-accepted">0</span></td></tr>
  </tbody>
  <tbody id="group-rejected">
    <tr class="group-header-row"><td colspan="12">&#x274C; Rejected <span class="group-count" id="count-rejected">0</span></td></tr>
  </tbody>
  <tbody id="group-staging" style="display:none">
{rows}
  </tbody>
</table>
</div>
<footer>
  Generated by the job-search agent engine. Decisions are stored in this browser, separately for each agent home,
  until you click "Sync decisions"; they reach the tracker on the next <code>./run.sh sync</code> or agent run.
  "New this week" means Date Found matches the most recent search run in this file.
</footer>
<script>
const REASONS = {reasons_json};
const HOME_ID = {home_id_js};
// Pages opened from local files share one browser storage in some browsers,
// so everything is kept per agent home.
const STORE_KEY = 'jobAgentDecisions:' + HOME_ID;

function loadState() {{
  try {{ return JSON.parse(localStorage.getItem(STORE_KEY)) || {{}}; }} catch (e) {{ return {{}}; }}
}}
function saveState(state) {{
  localStorage.setItem(STORE_KEY, JSON.stringify(state));
}}

function effectiveDecision(row) {{
  const key = row.getAttribute('data-key');
  const state = loadState();
  if (state[key]) return {{ source: 'pending', ...state[key] }};
  const initDecision = row.getAttribute('data-initial-decision');
  if (initDecision) return {{ source: 'synced', decision: initDecision, reason: row.getAttribute('data-initial-reason') || '' }};
  return null;
}}

function renderDecisionCell(row) {{
  const cell = row.querySelector('.decision-cell');
  const eff = effectiveDecision(row);
  const key = row.getAttribute('data-key');

  if (!eff) {{
    cell.innerHTML =
      '<div class="decision-buttons">' +
      '<button class="btn-accept" onclick="accept(this)">Accept</button>' +
      '<button class="btn-reject" onclick="showReasonPicker(this)">Reject</button>' +
      '</div>';
    return;
  }}

  let html = '';
  if (eff.decision === 'accepted') {{
    const drafted = row.getAttribute('data-drafted') === 'Y';
    html = drafted
      ? '<span class="decision-badge accepted">&check; Accepted &mdash; Documents generated</span>'
      : '<span class="decision-badge accepted">&check; Accepted &mdash; queued for drafting</span>';
  }} else {{
    const reasons = Array.isArray(eff.reasons) ? eff.reasons : (eff.reasons ? [eff.reasons] : (eff.reason ? [eff.reason] : []));
    html = '<span class="decision-badge rejected">&cross; Rejected<span class="decision-reason">' + escapeHtml(reasons.join('; ')) + '</span></span>';
  }}
  if (eff.source === 'pending') {{
    html += '<span class="pending-note">pending sync</span>';
    html += '<span class="undo-link" onclick="undo(\\'' + key.replace(/'/g, "\\\\'") + '\\')">Undo</span>';
  }}
  cell.innerHTML = html;
}}

function escapeHtml(s) {{
  const d = document.createElement('div');
  d.innerText = s;
  return d.innerHTML;
}}

function accept(btn) {{
  const row = btn.closest('tr');
  const key = row.getAttribute('data-key');
  const state = loadState();
  state[key] = {{ decision: 'accepted', ts: new Date().toISOString() }};
  saveState(state);
  renderDecisionCell(row);
  regroupRows();
  filterRows();
}}

function showReasonPicker(btn) {{
  const row = btn.closest('tr');
  const cell = row.querySelector('.decision-cell');
  let picker = '<div class="reason-picker">';
  picker += '<div class="r-title">Why reject this posting? (select all that apply)</div>';
  REASONS.forEach(function(r) {{
    picker += '<label><input type="checkbox" name="reason-' + row.rowIndex + '" value="' + escapeHtml(r) + '"> ' + escapeHtml(r) + '</label>';
  }});
  picker += '<input type="text" class="reject-note" placeholder="Optional note...">';
  picker += '<div class="r-actions"><button onclick="cancelReject(this)">Cancel</button><button class="confirm" onclick="confirmReject(this)">Confirm reject</button></div>';
  picker += '</div>';
  cell.innerHTML = picker;
}}

function cancelReject(btn) {{
  const row = btn.closest('tr');
  renderDecisionCell(row);
}}

function confirmReject(btn) {{
  const row = btn.closest('tr');
  const key = row.getAttribute('data-key');
  const picker = btn.closest('.reason-picker');
  const checked = picker.querySelectorAll('input[type=checkbox]:checked');
  if (checked.length === 0) {{
    alert('Select at least one reason.');
    return;
  }}
  const reasons = Array.from(checked).map(function(c) {{ return c.value; }});
  const note = picker.querySelector('.reject-note').value.trim();
  const state = loadState();
  state[key] = {{ decision: 'rejected', reasons: reasons, note: note, ts: new Date().toISOString() }};
  saveState(state);
  renderDecisionCell(row);
  regroupRows();
  filterRows();
}}

function undo(key) {{
  const state = loadState();
  delete state[key];
  saveState(state);
  document.querySelectorAll('#jobsTable tbody tr[data-key]').forEach(function(row) {{
    if (row.getAttribute('data-key') === key) renderDecisionCell(row);
  }});
  regroupRows();
  filterRows();
}}

const DIR_HANDLE_DB = 'jobAgentSyncDB';
const DIR_HANDLE_STORE = 'handles';
const DIR_HANDLE_KEY = 'state_dir:' + HOME_ID;
const HOME_ID_FILE = '.agent_home_id';

async function isThisHomesStateDir(dirHandle) {{
  try {{
    const f = await (await dirHandle.getFileHandle(HOME_ID_FILE)).getFile();
    return (await f.text()).trim() === HOME_ID;
  }} catch (e) {{ return false; }}
}}

function idbGet(key) {{
  return new Promise(function(resolve) {{
    const req = indexedDB.open(DIR_HANDLE_DB, 1);
    req.onupgradeneeded = function() {{ req.result.createObjectStore(DIR_HANDLE_STORE); }};
    req.onsuccess = function() {{
      const tx = req.result.transaction(DIR_HANDLE_STORE, 'readonly');
      const getReq = tx.objectStore(DIR_HANDLE_STORE).get(key);
      getReq.onsuccess = function() {{ resolve(getReq.result || null); }};
      getReq.onerror = function() {{ resolve(null); }};
    }};
    req.onerror = function() {{ resolve(null); }};
  }});
}}
function idbSet(key, value) {{
  return new Promise(function(resolve) {{
    const req = indexedDB.open(DIR_HANDLE_DB, 1);
    req.onupgradeneeded = function() {{ req.result.createObjectStore(DIR_HANDLE_STORE); }};
    req.onsuccess = function() {{
      const tx = req.result.transaction(DIR_HANDLE_STORE, 'readwrite');
      tx.objectStore(DIR_HANDLE_STORE).put(value, key);
      tx.oncomplete = function() {{ resolve(); }};
    }};
    req.onerror = function() {{ resolve(); }};
  }});
}}

function setSyncStatus(msg, isError) {{
  const el = document.getElementById('syncStatus');
  el.textContent = msg;
  el.style.color = isError ? 'var(--reject-fg)' : 'var(--accept-fg)';
  setTimeout(function() {{ el.textContent = ''; }}, 6000);
}}

function downloadFallback(state) {{
  const blob = new Blob([JSON.stringify(state, null, 2)], {{ type: 'application/json' }});
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a');
  const stamp = new Date().toISOString().slice(0, 10);
  a.href = url;
  a.download = 'decisions_sync_' + stamp + '.json';
  a.click();
  URL.revokeObjectURL(url);
  setSyncStatus('Downloaded - move it into {sync_hint_js} and rename it decisions_sync.json.');
}}

async function syncDecisions() {{
  const state = loadState();

  if (!window.showDirectoryPicker) {{
    downloadFallback(state);
    return;
  }}

  try {{
    let dirHandle = await idbGet(DIR_HANDLE_KEY);
    if (dirHandle) {{
      const perm = await dirHandle.queryPermission({{ mode: 'readwrite' }});
      if (perm !== 'granted') {{
        const req = await dirHandle.requestPermission({{ mode: 'readwrite' }});
        if (req !== 'granted') dirHandle = null;
      }}
    }}
    if (dirHandle && !(await isThisHomesStateDir(dirHandle))) dirHandle = null;
    if (!dirHandle) {{
      setSyncStatus('Pick {sync_hint_js} (one-time setup)...');
      dirHandle = await window.showDirectoryPicker({{ mode: 'readwrite' }});
      if (!(await isThisHomesStateDir(dirHandle))) {{
        setSyncStatus('That is not the state folder of this agent home ({sync_path_js}). Nothing was saved; click Sync again and pick that folder.', true);
        return;
      }}
      await idbSet(DIR_HANDLE_KEY, dirHandle);
    }}
    const fileHandle = await dirHandle.getFileHandle('decisions_sync.json', {{ create: true }});
    const writable = await fileHandle.createWritable();
    await writable.write(JSON.stringify(state, null, 2));
    await writable.close();
    setSyncStatus('Saved decisions_sync.json - applied on the next sync.');
  }} catch (err) {{
    if (err && err.name === 'AbortError') {{
      setSyncStatus('Cancelled.', true);
      return;
    }}
    downloadFallback(state);
  }}
}}

// Moves every row into the right section (new / undecided / accepted /
// rejected) based on its CURRENT effective decision, and updates the
// header counts. Runs on load and after every accept/reject/undo so a
// click moves the row immediately - no waiting for the next weekly
// regeneration or sync round-trip.
function regroupRows() {{
  const rows = Array.from(document.querySelectorAll('#jobsTable tbody tr[data-key]'));
  const counts = {{ new: 0, undecided: 0, accepted: 0, rejected: 0 }};
  rows.forEach(function(row) {{
    const eff = effectiveDecision(row);
    let group;
    if (eff && eff.decision === 'accepted') group = 'accepted';
    else if (eff && eff.decision === 'rejected') group = 'rejected';
    else if (row.getAttribute('data-new') === '1') group = 'new';
    else group = 'undecided';
    row.classList.toggle('is-new', row.getAttribute('data-new') === '1' && group !== 'accepted' && group !== 'rejected');
    counts[group]++;
    document.getElementById('group-' + group).appendChild(row);
  }});
  Object.keys(counts).forEach(function(g) {{
    const countEl = document.getElementById('count-' + g);
    if (countEl) countEl.textContent = counts[g];
    const headerRow = document.querySelector('#group-' + g + ' .group-header-row');
    if (headerRow) headerRow.classList.toggle('empty-group', counts[g] === 0);
  }});
  document.getElementById('newCount').textContent = counts.new;
  document.getElementById('acceptedCount').textContent = counts.accepted;
  document.getElementById('rejectedCount').textContent = counts.rejected;
}}

let currentTab = 'all';
function setTab(tab) {{
  currentTab = tab;
  document.querySelectorAll('.tab').forEach(function(t) {{
    t.classList.toggle('active', t.getAttribute('data-tab') === tab);
  }});
  filterRows();
}}

function filterRows() {{
  var q = document.getElementById('filterBox').value.toLowerCase();
  var onlyFlagged = document.getElementById('onlyFlagged').checked;
  var hideSenior = document.getElementById('hideSenior').checked;
  var rows = document.querySelectorAll('#jobsTable tbody tr[data-key]');
  rows.forEach(function(row) {{
    var text = row.getAttribute('data-search');
    var flagged = row.getAttribute('data-flagged') === '1';
    var senior = row.getAttribute('data-senior') === '1';
    var eff = effectiveDecision(row);
    var decisionState = eff ? eff.decision : 'undecided';
    var show = text.indexOf(q) !== -1;
    if (onlyFlagged && !flagged) show = false;
    if (hideSenior && senior) show = false;
    if (currentTab !== 'all' && decisionState !== currentTab) show = false;
    row.classList.toggle('hidden', !show);
  }});
}}

document.querySelectorAll('#jobsTable tbody tr[data-key]').forEach(renderDecisionCell);
regroupRows();
filterRows();
</script>
</body>
</html>
"""

ROW_TEMPLATE = """    <tr data-search="{search}" data-flagged="{flagged}" data-senior="{senior}" data-new="{is_new}" data-drafted="{drafted}" data-key="{key}" data-initial-decision="{init_decision}" data-initial-reason="{init_reason}">
      <td class="decision-cell decision-col"></td>
      <td class="employer">{employer}</td>
      <td>{title}{new_tag}</td>
      <td><span class="badge {sen_class}">{seniority}</span></td>
      <td>{location}</td>
      <td class="muted">{mode}</td>
      <td class="muted">{announce}</td>
      <td class="muted">{days}</td>
      <td class="muted">{category}</td>
      <td class="muted">{skills}</td>
      <td>{link}</td>
      <td class="muted">{status}</td>
    </tr>"""


def esc(v):
    return html.escape(str(v)) if v is not None else ""


SEN_CLASS = {"senior": "sen-senior", "junior": "sen-grad", "good": "sen-good", "other": "sen-other"}


def compute_action(deadline_val, announce_date, act_asap_after_days=7):
    today = date.today()
    d = to_date(deadline_val)
    if d is not None:
        return "⚠ ACT IMMEDIATELY" if 0 <= (d - today).days <= 1 else ""
    if isinstance(deadline_val, str) and "rolling" in deadline_val.lower() and announce_date is not None:
        days = (today - announce_date).days
        if days > act_asap_after_days:
            return f"⚠ ACT ASAP ({days}d since posting)"
    return ""


def initial_decision(status):
    s = str(status or "")
    if s.startswith("Approved"):
        return "accepted", ""
    if s.startswith("Rejected"):
        return "rejected", s[len("Rejected:"):].strip() if s.startswith("Rejected:") else ""
    return "", ""


def read_active_rows(path):
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Active Advertisements" not in wb.sheetnames:
        return []
    ws = wb["Active Advertisements"]
    headers = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(headers) if h}
        if "URL" in headers:
            cell = ws.cell(row=r, column=headers.index("URL") + 1)
            if cell.hyperlink:
                row["URL"] = cell.hyperlink.target
        if any(v not in (None, "") for k, v in row.items() if k not in ("Action", "Days Since Posting")):
            rows.append(row)
    return rows


def build_html(rows, rejection_reasons, seniority_cfg=None, weekly_note=None, generated_at=None,
               page_title="Job Search Tracker - Active Postings", sync_hint="the agent's state folder",
               act_asap_after_days=7, home_id="default", sync_path=""):
    today = date.today()
    found = [d for d in (to_date(r.get("Date Found")) for r in rows) if d]
    newest = max(found) if found else None
    out, urgent = [], 0
    for row in rows:
        announce = to_date(row.get("Announce Date"))
        action = compute_action(row.get("Deadline"), announce, act_asap_after_days)
        urgent += 1 if action else 0
        seniority = row.get("Seniority Level") or "Not stated"
        bucket = seniority_bucket(seniority, seniority_cfg)
        url = row.get("URL") or ""
        title = row.get("Title") or ""
        init_d, init_r = initial_decision(row.get("Status"))
        is_new = bool(newest and to_date(row.get("Date Found")) == newest)
        out.append(ROW_TEMPLATE.format(
            search=esc(" ".join(str(row.get(k, "") or "") for k in ["Employer", "Title", "Location", "Category", "Seniority Level"]).lower()),
            flagged="1" if action else "0",
            senior="1" if bucket == "senior" else "0",
            is_new="1" if is_new else "0",
            drafted=esc(str(row.get("CV/Cover Letter Drafted (Y/N)") or "N").strip().upper()),
            key=esc(f"{url}||{title}"),
            init_decision=esc(init_d), init_reason=esc(init_r),
            employer=esc(row.get("Employer")), title=esc(title),
            new_tag='<span class="new-tag">New</span>' if is_new else "",
            sen_class=SEN_CLASS[bucket], seniority=esc(seniority),
            location=esc(row.get("Location")), mode=esc(row.get("Remote/Hybrid/Onsite")),
            announce=esc(announce.isoformat() if announce else row.get("Announce Date") or ""),
            days=(today - announce).days if announce else "",
            category=esc(row.get("Category")), skills=esc(row.get("Required Skills (summary)")),
            link=f'<a class="posting-link" href="{esc(url)}" target="_blank" rel="noopener">View posting &rarr;</a>' if url else "",
            status=esc(row.get("Status")),
        ))
    note_html = (f'<span class="muted" style="display:block;font-size:11px;font-style:italic;margin-top:2px;">{esc(weekly_note)}</span>'
                 if weekly_note else "")
    return PAGE_TEMPLATE.format(
        page_title=esc(page_title), page_title_html=esc(page_title).replace(" - ", " &mdash; "),
        active_count=len(rows), urgent_count=urgent,
        generated_at=generated_at or today.isoformat(), weekly_note_html=note_html,
        sync_hint_attr=esc(sync_hint), sync_hint_js=sync_hint.replace("\\", "\\\\").replace("'", "\\'"),
        home_id_js=json.dumps(home_id), sync_path_js=sync_path.replace("\\", "\\\\").replace("'", "\\'"),
        rows="\n".join(out), reasons_json=json.dumps(rejection_reasons),
    )


def home_id(home):
    """Random, stable ID of an agent home, kept in <state folder>/.agent_home_id.
    The weekly page uses it to keep its browser storage per home and to
    check that Sync writes into the right folder."""
    import uuid
    p = os.path.join(os.path.dirname(home.path("decisions")), ".agent_home_id")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            v = f.read().strip()
        if v:
            return v
    os.makedirs(os.path.dirname(p), exist_ok=True)
    v = uuid.uuid4().hex
    with open(p, "w", encoding="utf-8") as f:
        f.write(v + "\n")
    return v


def generate(home, weekly_note=None, out_path=None):
    cfg = home.search_config()
    rows = read_active_rows(home.path("tracker"))
    out_path = out_path or home.path("weekly_html")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    sync_dir = os.path.dirname(home.path("decisions"))
    rel = os.path.relpath(sync_dir, home.root)
    hint = "your agent home folder" if rel == "." else f"the '{rel}' folder in your agent home"
    html_text = build_html(rows, cfg.get("rejection_reasons") or [], seniority_cfg=cfg.get("seniority"),
                           home_id=home_id(home), sync_path=sync_dir,
                           weekly_note=weekly_note, generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
                           page_title=cfg.get("page_title") or "Job Search Tracker - Active Postings",
                           sync_hint=hint, act_asap_after_days=int(cfg.get("act_asap_after_days", 7)))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html_text)
    print(f"wrote {out_path} ({len(rows)} active postings)")
    return out_path
