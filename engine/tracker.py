# Job-search agent engine
# File:    engine/tracker.py
# Purpose: Reads and rewrites the tracker workbook without ever dropping a row.
# Author:  manpmanp <manp.node@gmail.com>
# Generated with Claude (Anthropic), under the author's direction and review.
# License: MIT (see LICENSE)
# Version: 0.2.0
# Date:    2026-10-28

"""
The tracker workbook (Job_Search_Tracker.xlsx): the complete record of every
posting the agent has found.

Three sheets together hold every row; nothing is ever deleted, only moved:
- Active Advertisements: everything not Expired and not Rejected, with a
  live urgency formula, seniority colour-coding and collapsible detail
  columns (plain Excel features, no macros).
- Rejected Advertisements: Status starts with "Rejected".
- Expired Advertisements: Status == "Expired".

update() is idempotent: it reads every existing row, merges the rows you
pass in (yours win for the same (URL, Title) key), de-duplicates by
(URL, Title), sorts newest-first by Date Found and rewrites all sheets.
A run with no new rows never drops a row.
"""
import os
import re
from datetime import datetime, date

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

ACTIVE_LAYOUT = ["Action", "Employer", "Title", "Seniority Level", "Location", "Remote/Hybrid/Onsite",
                 "Announce Date", "Days Since Posting", "Date Found", "Category", "Deadline",
                 "Required Skills (summary)", "URL", "Calendar Event Added (Y/N)",
                 "CV/Cover Letter Drafted (Y/N)", "Status", "Notes"]
FLAT_LAYOUT = ["Date Found", "Title", "Employer", "Location", "Category", "Seniority Level",
               "Announce Date", "Deadline", "Remote/Hybrid/Onsite", "Required Skills (summary)",
               "URL", "Calendar Event Added (Y/N)", "CV/Cover Letter Drafted (Y/N)", "Status", "Notes"]
ROW_FIELDS = FLAT_LAYOUT

ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

DEFAULT_SENIORITY = {
    "flag_as_too_senior": ["Senior", "Staff", "Lead", "Principal", "Architect", "architect"],
    "flag_as_too_junior": ["Graduate", "Entry"],
    "preferred": ["Experienced", "Mid", "mid"],
}


def to_date(val):
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str) and ISO_DATE_RE.match(val.strip()):
        try:
            return datetime.strptime(val.strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def seniority_bucket(value, seniority_cfg=None):
    """'senior' | 'junior' | 'good' | 'other' using the configured lists."""
    cfg = dict(DEFAULT_SENIORITY)
    cfg.update({k: v for k, v in (seniority_cfg or {}).items() if v})
    # Case-sensitive on purpose: "Senior" is a level label, while
    # "Experienced (mid-senior)" or "leads a team" are not senior titles.
    v = str(value or "")
    if any(k in v for k in cfg["flag_as_too_senior"]):
        return "senior"
    if any(k in v for k in cfg["flag_as_too_junior"]):
        return "junior"
    if any(k in v for k in cfg["preferred"]):
        return "good"
    return "other"


def _label(terms):
    seen, out = set(), []
    for t in terms:
        if t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
    return " / ".join(out)


def status_bucket(status):
    s = str(status or "")
    if s == "Expired":
        return "expired"
    if s.startswith("Rejected"):
        return "rejected"
    return "active"


def _read_sheet(ws, skip=("Action",)):
    headers = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(headers) if h and h not in skip}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return rows


def read_all_rows(path):
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path)
    rows = []
    for name in ("Active Advertisements", "Rejected Advertisements", "Expired Advertisements"):
        if name in wb.sheetnames:
            rows += _read_sheet(wb[name])
    return rows


def _style_sheet(ws, n_cols, n_rows, freeze_cell="A2"):
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    light_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=1, max_row=n_rows, max_col=n_cols):
        for cell in row:
            cell.border = border
    for r in range(2, n_rows + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if r % 2 == 0:
            for c in range(1, n_cols + 1):
                rgb = ws.cell(row=r, column=c).fill.start_color.rgb
                if not rgb or rgb == "00000000":
                    ws.cell(row=r, column=c).fill = light_fill
        ws.row_dimensions[r].height = 46
    ws.freeze_panes = freeze_cell
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{max(n_rows, 1)}"


def update(path, new_rows=None, seniority_cfg=None, act_asap_after_days=7):
    new_rows = [dict(r) for r in (new_rows or [])]
    for r in new_rows:
        r.setdefault("Calendar Event Added (Y/N)", "N")
        r.setdefault("CV/Cover Letter Drafted (Y/N)", "N")
        r.setdefault("Status", "New")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = openpyxl.load_workbook(path) if os.path.exists(path) else openpyxl.Workbook()

    existing = []
    for name in ("Active Advertisements", "Rejected Advertisements", "Expired Advertisements"):
        if name in wb.sheetnames:
            existing += _read_sheet(wb[name])

    # Passed-in rows come first so they win the dedup (first seen per key).
    seen, deduped = set(), []
    for r in new_rows + existing:
        key = (r.get("URL"), r.get("Title"))
        if key in seen:
            continue
        seen.add(key)
        deduped.append(r)

    def _sk(r):
        d = to_date(r.get("Date Found")) or date(1900, 1, 1)
        return datetime(d.year, d.month, d.day)
    deduped.sort(key=_sk, reverse=True)

    active = [r for r in deduped if status_bucket(r.get("Status")) == "active"]
    rejected = [r for r in deduped if status_bucket(r.get("Status")) == "rejected"]
    expired = [r for r in deduped if status_bucket(r.get("Status")) == "expired"]

    for name in ["Active Advertisements", "Rejected Advertisements", "Expired Advertisements", "Legend", "Tracker", "Sheet"]:
        if name in wb.sheetnames:
            del wb[name]

    link_font = dict(name="Arial", size=10, color="0563C1", underline="single")

    # ---------------- Active ----------------
    ws_a = wb.create_sheet("Active Advertisements", 0)
    for c, h in enumerate(ACTIVE_LAYOUT, start=1):
        ws_a.cell(row=1, column=c, value=h)
    deadline_col = ACTIVE_LAYOUT.index("Deadline") + 1
    announce_col = ACTIVE_LAYOUT.index("Announce Date") + 1
    days_col = ACTIVE_LAYOUT.index("Days Since Posting") + 1
    senior_col = ACTIVE_LAYOUT.index("Seniority Level") + 1
    url_col = ACTIVE_LAYOUT.index("URL") + 1
    dl = get_column_letter(deadline_col)
    al = get_column_letter(announce_col)

    for i, row in enumerate(active, start=2):
        vals = []
        for h in ACTIVE_LAYOUT:
            if h in ("Action", "Days Since Posting"):
                vals.append(None)
            elif h in ("Announce Date", "Deadline"):
                vals.append(to_date(row.get(h)) or row.get(h))
            else:
                vals.append(row.get(h))
        for c, v in enumerate(vals, start=1):
            cell = ws_a.cell(row=i, column=c, value=v)
            if c in (announce_col, deadline_col) and isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
            if c == url_col and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(**link_font)
        ws_a.cell(row=i, column=days_col, value=f'=IF(ISNUMBER({al}{i}),TODAY()-{al}{i},"")')
        ws_a.cell(row=i, column=1, value=(
            f'=IF(ISNUMBER({dl}{i}),'
            f'IF(AND(({dl}{i}-TODAY())>=0,({dl}{i}-TODAY())<=1),"⚠ ACT IMMEDIATELY",""),'
            f'IF(ISNUMBER(SEARCH("rolling",{dl}{i})),'
            f'IF(AND(ISNUMBER({al}{i}),(TODAY()-{al}{i})>{int(act_asap_after_days)}),'
            f'"⚠ ACT ASAP ("&(TODAY()-{al}{i})&"d since posting)","")'
            f',""))'))

    n_active = len(active) + 1
    _style_sheet(ws_a, len(ACTIVE_LAYOUT), n_active, freeze_cell="I2")
    widths = {
        "Action": 22, "Employer": 18, "Title": 30, "Seniority Level": 22, "Location": 16,
        "Remote/Hybrid/Onsite": 14, "Announce Date": 13, "Days Since Posting": 13, "Date Found": 11,
        "Category": 24, "Deadline": 13, "Required Skills (summary)": 42, "URL": 40,
        "Calendar Event Added (Y/N)": 10, "CV/Cover Letter Drafted (Y/N)": 10, "Status": 15, "Notes": 42,
    }
    for i, h in enumerate(ACTIVE_LAYOUT, start=1):
        ws_a.column_dimensions[get_column_letter(i)].width = widths.get(h, 18)

    red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    red_font = Font(name="Arial", size=10, bold=True, color="842029")
    ws_a.conditional_formatting.add(f"A2:A{max(n_active, 2)}", FormulaRule(formula=['LEN($A2)>0'], fill=red_fill, font=red_font))

    fills = {
        "senior": (PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid"), Font(name="Arial", size=10, color="9C5700", bold=True)),
        "junior": (PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"), Font(name="Arial", size=10, color="842029", bold=True)),
        "good": (PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), Font(name="Arial", size=10, color="155724", bold=True)),
    }
    for r in range(2, n_active + 1):
        cell = ws_a.cell(row=r, column=senior_col)
        b = seniority_bucket(cell.value, seniority_cfg)
        if b in fills:
            cell.fill, cell.font = fills[b]

    for i in range(9, len(ACTIVE_LAYOUT) + 1):
        cd = ws_a.column_dimensions[get_column_letter(i)]
        cd.outlineLevel = 1
        cd.hidden = True
    ws_a.sheet_properties.outlinePr.summaryRight = True

    # ---------------- Flat sheets ----------------
    flat_widths = {
        "Date Found": 11, "Title": 30, "Employer": 18, "Location": 18, "Category": 24,
        "Seniority Level": 22, "Announce Date": 14, "Deadline": 18, "Remote/Hybrid/Onsite": 14,
        "Required Skills (summary)": 42, "URL": 40, "Calendar Event Added (Y/N)": 10,
        "CV/Cover Letter Drafted (Y/N)": 10, "Status": 13, "Notes": 42,
    }

    def flat_sheet(name, index, rows, fill_hex, font_color):
        ws = wb.create_sheet(name, index)
        for c, h in enumerate(FLAT_LAYOUT, start=1):
            ws.cell(row=1, column=c, value=h)
        uc = FLAT_LAYOUT.index("URL") + 1
        for i, row in enumerate(rows, start=2):
            for c, h in enumerate(FLAT_LAYOUT, start=1):
                v = row.get(h)
                cell = ws.cell(row=i, column=c, value=v)
                if c == uc and isinstance(v, str) and v.startswith("http"):
                    cell.hyperlink = v
        n = len(rows) + 1
        _style_sheet(ws, len(FLAT_LAYOUT), n, freeze_cell="A2")
        for i, h in enumerate(FLAT_LAYOUT, start=1):
            ws.column_dimensions[get_column_letter(i)].width = flat_widths.get(h, 18)
        f = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=n, max_col=len(FLAT_LAYOUT)):
            for cell in row:
                cell.fill = f
                cell.font = Font(name="Arial", size=10, italic=True, color=font_color)

    flat_sheet("Rejected Advertisements", 1, rejected, "FBE4E4", "842029")
    flat_sheet("Expired Advertisements", 2, expired, "E2E2E2", "6B6B6B")

    legend = wb.create_sheet("Legend")
    legend["A1"] = "Seniority colour legend (from search_config.seniority)"
    legend["A1"].font = Font(name="Arial", size=12, bold=True)
    cfg = dict(DEFAULT_SENIORITY)
    cfg.update({k: v for k, v in (seniority_cfg or {}).items() if v})
    for r, (label, key, note) in enumerate([
        (_label(cfg["flag_as_too_senior"]), "senior", "Flagged as likely too senior - review before applying"),
        (_label(cfg["preferred"]), "good", "Preferred seniority"),
        (_label(cfg["flag_as_too_junior"]), "junior", "Flagged as likely too junior"),
    ], start=3):
        c1 = legend.cell(row=r, column=1, value=label)
        c1.fill, c1.font = fills[key]
        c2 = legend.cell(row=r, column=2, value=note)
        c2.font = Font(name="Arial", size=10)
    legend.column_dimensions["A"].width = 40
    legend.column_dimensions["B"].width = 55

    wb._sheets = [wb[n] for n in ["Active Advertisements", "Rejected Advertisements", "Expired Advertisements", "Legend"]]
    wb.save(path)
    counts = {"active": len(active), "rejected": len(rejected), "expired": len(expired), "total": len(deduped)}
    print(f"tracker updated - active: {counts['active']}, rejected: {counts['rejected']}, expired: {counts['expired']}, total: {counts['total']}")
    return counts
